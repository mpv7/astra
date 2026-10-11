"""Генерация сводного xlsx-отчёта (Задание 2.2, п.3)."""
from datetime import datetime
from io import BytesIO
from typing import Any

from django.db.models import QuerySet
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .aggregation import aggregate, macro_by_period

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
METRICS = ["count", "sum_quantity", "sum_revenue", "avg_unit_price", "avg_discount_pct"]
METRIC_TITLES = {"count": "Транзакций", "sum_quantity": "Количество", "sum_revenue": "Выручка, $",
                 "avg_unit_price": "Средняя цена, $", "avg_discount_pct": "Средняя скидка, %"}
NUMBER_FORMATS = {"count": "#,##0", "sum_quantity": "#,##0", "sum_revenue": "#,##0.00",
                  "avg_unit_price": "#,##0.00", "avg_discount_pct": "0.00", "share": "0.0%"}


def _write_table(ws, headers: list[str], rows: list[list[Any]], formats: list[str | None], start_row: int = 1) -> None:
    """Пишет таблицу со стилизованной шапкой, форматами, автошириной и автофильтром."""
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            if formats[c - 1]:
                cell.number_format = formats[c - 1]
    for c in range(1, len(headers) + 1):
        longest = max([len(str(headers[c - 1]))] + [len(str(row[c - 1])) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(c)].width = min(max(12, longest + 2), 40)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def _dimension_sheet(wb: Workbook, queryset: QuerySet, dim: str, title: str, chart: bool = False) -> None:
    """Лист с разрезом по категориальному полю: метрики + доля выручки (+ столбчатая диаграмма)."""
    rows = sorted(aggregate(queryset, [dim], METRICS), key=lambda r: -(r["sum_revenue"] or 0))
    total = sum(r["sum_revenue"] or 0 for r in rows) or 1
    ws = wb.create_sheet(title)
    headers = [title] + [METRIC_TITLES[m] for m in METRICS] + ["Доля выручки"]
    data = [[r[dim] or "—"] + [r[m] for m in METRICS] + [(r["sum_revenue"] or 0) / total] for r in rows]
    _write_table(ws, headers, data, [None] + [NUMBER_FORMATS[m] for m in METRICS] + [NUMBER_FORMATS["share"]])
    if chart and rows:
        bar = BarChart()
        bar.title, bar.y_axis.title, bar.height, bar.width = f"Выручка по: {title.lower()}", "$", 9, 20
        bar.add_data(Reference(ws, min_col=4, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1))
        bar.legend = None
        ws.add_chart(bar, f"{get_column_letter(len(headers) + 2)}2")


def build_summary_workbook(queryset: QuerySet, filters_note: str = "") -> bytes:
    """Строит сводный отчёт по отфильтрованным продажам и возвращает содержимое xlsx.

    Листы: «Сводка», «По месяцам» (с макропоказателями и графиком), «Категории», «Регионы»,
    «Каналы», «Сегменты».
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    totals = aggregate(queryset, [], METRICS + ["sum_cost"])[0]
    ws["A1"] = "Сводный отчёт по продажам строительных материалов"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Сформирован: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A3"] = f"Фильтры: {filters_note or 'не заданы (все данные)'}"
    kpis = [(METRIC_TITLES[m], totals[m], NUMBER_FORMATS[m]) for m in METRICS]
    if totals.get("sum_cost") and totals.get("sum_revenue"):
        kpis.append(("Валовая маржа", (totals["sum_revenue"] - totals["sum_cost"]) / totals["sum_revenue"], "0.0%"))
    for i, (name, value, fmt) in enumerate(kpis, 5):
        ws.cell(row=i, column=1, value=name).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value).number_format = fmt
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 22

    macro = macro_by_period("month")
    indicators = sorted({name for values in macro.values() for name in values})
    months = aggregate(queryset, ["month"], METRICS)
    ws_m = wb.create_sheet("По месяцам")
    headers = ["Месяц"] + [METRIC_TITLES[m] for m in METRICS] + [f"Макро: {n}" for n in indicators]
    data = [[r["month"], *[r[m] for m in METRICS], *[macro.get(r["month"], {}).get(n) for n in indicators]] for r in months]
    _write_table(ws_m, headers, data,
                 ["yyyy-mm"] + [NUMBER_FORMATS[m] for m in METRICS] + ["0.00"] * len(indicators))
    if months:
        line = LineChart()
        line.title, line.height, line.width = "Выручка по месяцам", 9, 22
        line.add_data(Reference(ws_m, min_col=4, min_row=1, max_row=len(months) + 1), titles_from_data=True)
        line.set_categories(Reference(ws_m, min_col=1, min_row=2, max_row=len(months) + 1))
        line.legend = None
        ws_m.add_chart(line, f"{get_column_letter(len(headers) + 2)}2")

    _dimension_sheet(wb, queryset, "category", "Категории", chart=True)
    _dimension_sheet(wb, queryset, "region", "Регионы", chart=True)
    _dimension_sheet(wb, queryset, "sales_channel", "Каналы")
    _dimension_sheet(wb, queryset, "customer_segment", "Сегменты")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
