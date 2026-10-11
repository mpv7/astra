"""Тесты: импорт, API (фильтры, агрегации, макропоказатели), отчёт. Запуск: DB_ENGINE=sqlite python manage.py test"""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from .importer import parse_date, parse_decimal, record_to_sales, resolve_mapping
from .models import MacroIndicator, SalesTransaction
from .services.reports import build_summary_workbook
from .utils import week_start_of


class ImporterTests(TestCase):
    """Проверки разбора значений и сопоставления колонок."""

    def test_parse_decimal(self):
        self.assertEqual(parse_decimal("$1,234.50"), Decimal("1234.50"))
        self.assertIsNone(parse_decimal(""))
        self.assertIsNone(parse_decimal("abc"))

    def test_parse_date_formats(self):
        self.assertEqual(parse_date("2023-05-17"), date(2023, 5, 17))
        self.assertEqual(parse_date("05/17/2023"), date(2023, 5, 17))
        with self.assertRaises(ValueError):
            parse_date("не дата")

    def test_week_start_is_monday(self):
        self.assertEqual(week_start_of(date(2023, 5, 17)), date(2023, 5, 15))

    def test_mapping_and_extra(self):
        mapping = resolve_mapping(["week", "product_category", "units", "price", "note"])
        self.assertEqual(mapping["date"], "week")
        self.assertEqual(mapping["category"], "product_category")
        obj = record_to_sales({"week": "2023-05-15", "product_category": "Lumber", "units": "10",
                               "price": "5", "note": "x"}, mapping)
        self.assertEqual(obj.revenue, Decimal("50.00"))
        self.assertEqual(obj.extra, {"note": "x"})


class ApiTests(TestCase):
    """Проверки REST API и отчёта на небольшом наборе данных."""

    @classmethod
    def setUpTestData(cls):
        rows = [("West", "Lumber", "2023-05-15", 100), ("West", "Roofing", "2023-05-22", 50),
                ("South", "Lumber", "2024-01-08", 200)]
        for region, category, day, revenue in rows:
            SalesTransaction.objects.create(date=date.fromisoformat(day), region=region, category=category,
                                            quantity=1, unit_price=revenue, revenue=revenue)
        MacroIndicator.objects.create(week_start=date(2023, 5, 15), indicator="mortgage_rate", value=6.5)

    def test_filters(self):
        data = self.client.get("/api/sales/?region=West&year=2023").json()
        self.assertEqual(data["count"], 2)
        data = self.client.get("/api/sales/?category=Lumber,Roofing&revenue_min=100").json()
        self.assertEqual(data["count"], 2)

    def test_include_macro(self):
        results = self.client.get("/api/sales/?year=2023&include_macro=true&ordering=week_start").json()["results"]
        self.assertEqual(results[0]["macro"], {"mortgage_rate": 6.5})
        self.assertEqual(results[1]["macro"], {})

    def test_aggregate(self):
        data = self.client.get("/api/sales/aggregate/?group_by=category&metrics=sum_revenue").json()
        totals = {r["category"]: r["sum_revenue"] for r in data["results"]}
        self.assertEqual(totals, {"Lumber": 300.0, "Roofing": 50.0})

    def test_aggregate_validation(self):
        self.assertEqual(self.client.get("/api/sales/aggregate/?group_by=nonsense").status_code, 400)

    def test_report_is_valid_xlsx(self):
        wb = load_workbook(BytesIO(build_summary_workbook(SalesTransaction.objects.all())))
        self.assertIn("Сводка", wb.sheetnames)
        self.assertIn("Категории", wb.sheetnames)
