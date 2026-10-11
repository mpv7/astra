"""Чтение CSV/XLSX и приведение строк источника к моделям.

Схема Kaggle-датасета может отличаться от ожидаемой, поэтому колонки сопоставляются
по именам-синонимам (``COLUMN_ALIASES``). Сопоставление можно переопределить флагом
``--map поле=колонка`` у команды ``load_sales``. Незамеченные колонки уходят в ``extra``.
"""
import csv
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Iterator

from .models import SalesTransaction
from .utils import week_start_of

COLUMN_ALIASES: dict[str, list[str]] = {
    "transaction_id": ["transaction_id", "transaction", "txn_id", "order_id", "invoice_id", "id"],
    "date": ["date", "week", "week_start", "week_date", "week_ending", "week_end", "transaction_date",
             "order_date", "sale_date", "sales_date", "period"],
    "region": ["region", "sales_region", "census_region"],
    "state": ["state", "state_code", "us_state"],
    "category": ["category", "product_category", "material_category", "material_type", "category_name"],
    "product": ["product", "product_name", "item", "item_name", "sku", "material", "product_id"],
    "brand": ["brand", "supplier", "vendor", "manufacturer"],
    "customer_segment": ["customer_segment", "customer_type", "segment", "customer"],
    "sales_channel": ["sales_channel", "channel", "distribution_channel"],
    "quantity": ["quantity", "qty", "units", "units_sold", "volume"],
    "unit_price": ["unit_price", "price", "price_per_unit", "avg_price", "average_price"],
    "discount_pct": ["discount_pct", "discount", "discount_percent", "discount_rate"],
    "revenue": ["revenue", "sales", "sales_amount", "total_sales", "total_revenue", "amount", "total", "net_revenue"],
    "cost": ["cost", "cogs", "total_cost", "unit_cost"],
}
DATE_ALIASES = COLUMN_ALIASES["date"] + ["observation_date", "date_week", "weekly_date"]
TEXT_LIMITS = {"transaction_id": 64, "region": 64, "state": 64, "category": 128, "product": 255,
               "brand": 128, "customer_segment": 64, "sales_channel": 64}
DECIMAL_FIELDS = ["quantity", "unit_price", "discount_pct", "revenue", "cost"]
DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%y", "%d/%m/%Y")


def normalize_header(header: Any) -> str:
    """Приводит заголовок колонки к snake_case в нижнем регистре."""
    return re.sub(r"[^0-9a-z]+", "_", str(header).strip().lower()).strip("_")


def list_sheets(path: Path) -> list[str]:
    """Возвращает имена листов xlsx-файла (для CSV — пустой список)."""
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def iter_records(path: Path, sheet: str | None = None, delimiter: str = ",") -> Iterator[dict[str, Any]]:
    """Построчно читает CSV или XLSX и отдаёт словари с нормализованными заголовками."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            header = [normalize_header(c) if c is not None else f"col_{i}" for i, c in enumerate(next(rows))]
            for row in rows:
                if any(v is not None for v in row):
                    yield dict(zip(header, row))
        finally:
            wb.close()
        return
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        header = [normalize_header(h) for h in next(reader)]
        for row in reader:
            if row:
                yield dict(zip(header, row))


def resolve_mapping(headers: Iterable[str], overrides: dict[str, str] | None = None) -> dict[str, str]:
    """Сопоставляет поля модели с колонками файла: сначала явные ``overrides``, затем синонимы."""
    available = list(headers)
    mapping: dict[str, str] = {}
    taken: set[str] = set()
    for field, column in (overrides or {}).items():
        mapping[field] = normalize_header(column)
        taken.add(mapping[field])
    for field, aliases in COLUMN_ALIASES.items():
        if field in mapping:
            continue
        for alias in aliases:
            if alias in available and alias not in taken:
                mapping[field] = alias
                taken.add(alias)
                break
    return mapping


def parse_decimal(value: Any) -> Decimal | None:
    """Преобразует значение ячейки в Decimal; пустые и нечисловые значения дают None."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace("$", "").replace(",", "").replace("%", "").replace(" ", "")
    if text == "" or text.lower() in {"nan", "null", "none", "n/a", "-"}:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def parse_date(value: Any) -> date:
    """Разбирает дату из datetime/date или строки в одном из известных форматов."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Не удалось разобрать дату: {value!r}")


def _json_safe(value: Any) -> Any:
    """Делает значение пригодным для JSONField."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    json.dumps(value)
    return value


def _fallback_revenue(quantity: Decimal | None, price: Decimal | None, discount: Decimal | None) -> Decimal | None:
    """Считает выручку как количество × цена × (1 − скидка), если колонки revenue нет."""
    if quantity is None or price is None:
        return None
    rate = Decimal(0)
    if discount is not None:
        rate = discount / 100 if discount > 1 else discount
    return (quantity * price * (1 - rate)).quantize(Decimal("0.01"))


def record_to_sales(record: dict[str, Any], mapping: dict[str, str]) -> SalesTransaction:
    """Строит (несохранённый) объект SalesTransaction из строки источника.

    Raises:
        ValueError: если дата отсутствует или не разбирается — такая строка пропускается.
    """
    raw_date = record.get(mapping["date"])
    if raw_date in (None, ""):
        raise ValueError("пустая дата")
    day = parse_date(raw_date)
    data: dict[str, Any] = {"date": day, "week_start": week_start_of(day)}
    for field, limit in TEXT_LIMITS.items():
        column = mapping.get(field)
        value = record.get(column) if column else None
        data[field] = "" if value is None else str(value).strip()[:limit]
    for field in DECIMAL_FIELDS:
        column = mapping.get(field)
        data[field] = parse_decimal(record.get(column)) if column else None
    if data["revenue"] is None:
        data["revenue"] = _fallback_revenue(data["quantity"], data["unit_price"], data["discount_pct"])
    used = set(mapping.values())
    data["extra"] = {k: _json_safe(v) for k, v in record.items() if k not in used and v not in (None, "")}
    return SalesTransaction(**data)
